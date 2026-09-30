import io
import logging
import os
from dataclasses import dataclass
from typing import Any

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    content: str
    metadata: dict[str, Any]
    doc_type: str


def parse_pdf(file_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        import fitz
        import pymupdf4llm
    except ImportError:
        raise ImportError(
            "pymupdf4llm and pymupdf required for PDF parsing. Install with: pip install pymupdf4llm pymupdf"
        )

    try:
        markdown_text = pymupdf4llm.to_markdown(io.BytesIO(file_bytes))
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        metadata = {
            "source": filename,
            "doc_type": "pdf",
            "page_count": doc.page_count,
            "title": doc.metadata.get("title", filename),
            "author": doc.metadata.get("author", ""),
        }
        doc.close()
        return ParsedDocument(content=markdown_text, metadata=metadata, doc_type="pdf")
    except Exception as e:
        logger.error(f"PDF parsing failed for {filename}: {e}")
        raise


def parse_docx(file_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ImportError(
            "python-docx required for DOCX parsing. Install with: pip install python-docx"
        )

    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    full_text.append(row_text)

        content = "\n\n".join(full_text)
        metadata = {
            "source": filename,
            "doc_type": "docx",
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        }
        return ParsedDocument(content=content, metadata=metadata, doc_type="docx")
    except Exception as e:
        logger.error(f"DOCX parsing failed for {filename}: {e}")
        raise


def parse_csv(file_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        import pandas as pd
    except ImportError:
        raise ImportError(
            "pandas required for CSV parsing. Install with: pip install pandas"
        )

    try:
        df = pd.read_csv(io.BytesIO(file_bytes))
        markdown_table = df.to_markdown(index=False)
        content = f"# {filename}\n\n{markdown_table}"
        metadata = {
            "source": filename,
            "doc_type": "csv",
            "row_count": len(df),
            "column_count": len(df.columns),
            "columns": list(df.columns),
        }
        return ParsedDocument(content=content, metadata=metadata, doc_type="csv")
    except Exception as e:
        logger.error(f"CSV parsing failed for {filename}: {e}")
        raise


def parse_markdown(file_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        content = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = file_bytes.decode("latin-1")

    metadata = {
        "source": filename,
        "doc_type": "md",
        "char_count": len(content),
    }
    return ParsedDocument(content=content, metadata=metadata, doc_type="md")


def parse_text(file_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        content = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = file_bytes.decode("latin-1")

    metadata = {
        "source": filename,
        "doc_type": "txt",
        "char_count": len(content),
    }
    return ParsedDocument(content=content, metadata=metadata, doc_type="txt")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".csv": parse_csv,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".txt": parse_text,
}


def parse_document(file_bytes: bytes, filename: str) -> ParsedDocument:
    _, ext = os.path.splitext(filename.lower())
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(
            f"Unsupported file type: {ext}. Supported: {list(PARSERS.keys())}"
        )
    return parser(file_bytes, filename)


def get_supported_extensions() -> list[str]:
    return list(PARSERS.keys())
